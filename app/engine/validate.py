"""Validate a decision packet against a catalog, plus optional local file checks.

``validate_packet`` does not open raw corpus files. DOCX, PDF, and manifest
checks are separate and are meant for a machine that has the rendered outputs.
"""

from __future__ import annotations

import json
import re
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree

from docx import Document
from pypdf import PdfReader

from .render import (
    BLOCK_MARKERS,
    IGNORED_PARAGRAPH_KEYS,
    PUBLIC_BASE_URL,
    PUBLIC_FOLDERS_URL,
    WARNING,
    _collect_sources,
    public_document_url,
    repository_path,
)


ALLOWED_STATUSES = {"proved", "not_proved", "procedural"}
ALLOWED_BASES = {
    "common",
    "admission",
    "procedural_record",
    "unilateral_assertion",
    "unseen_attachment",
}
STATUS_BASES = {
    "proved": {"common", "admission", "procedural_record"},
    "procedural": {"procedural_record"},
    "not_proved": {"unilateral_assertion", "unseen_attachment"},
}
REQUIRED_HEADINGS = (
    "SUMÁRIO",
    "I – RELATÓRIO",
    "II – OBJETO DO LITÍGIO",
    "III – SANEADOR",
    "IV – FUNDAMENTAÇÃO",
    "V – MOTIVAÇÃO",
    "VI – FUNDAMENTAÇÃO DE DIREITO",
    "VII – DECISÃO",
)
FORBIDDEN_IDENTITY_FIELDS = {
    "judge",
    "judge_name",
    "arbitrator",
    "arbitrator_name",
    "signature",
    "signature_name",
}
OFFICE_RELATIONSHIPS_NAMESPACE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
)
WORD_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
SHA256_PATTERN = re.compile(r"^[0-9a-f]{64}$")


def _load_json(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ValueError(f"Cannot read JSON {path}: {exc}") from exc
    if not isinstance(value, dict):
        raise ValueError(f"JSON root must be an object: {path}")
    return value


def _paragraph_text(spec: Any) -> str:
    if isinstance(spec, str):
        return spec
    if isinstance(spec, dict) and isinstance(spec.get("text"), str):
        return spec["text"]
    return ""


def _iter_packet_text(value: Any) -> Iterable[str]:
    if isinstance(value, str):
        yield value
    elif isinstance(value, dict):
        paragraph = isinstance(value.get("text"), str)
        for key, item in value.items():
            if paragraph and key in IGNORED_PARAGRAPH_KEYS:
                continue
            yield from _iter_packet_text(item)
    elif isinstance(value, list):
        for item in value:
            yield from _iter_packet_text(item)


def _validate_metadata(packet: dict[str, Any], errors: list[str]) -> None:
    metadata = packet.get("metadata")
    required = (
        "process_number",
        "claimant",
        "respondent",
        "case_value",
        "place_and_date",
        "draft_date",
        "document_title",
    )
    if not isinstance(metadata, dict):
        errors.append("metadata must be an object")
        return
    for key in required:
        if not isinstance(metadata.get(key), str) or not metadata[key].strip():
            errors.append(f"metadata.{key} must be non-empty text")
    for field in FORBIDDEN_IDENTITY_FIELDS:
        if field in metadata and metadata.get(field) is not None and metadata.get(field) != "":
            errors.append(f"metadata must not identify or sign for an arbitrator: {field}")


def _validate_policy(packet: dict[str, Any], errors: list[str]) -> None:
    policy = packet.get("policy")
    if not isinstance(policy, dict):
        errors.append("policy must be an object")
        return
    if policy.get("strict_evidence") is not True:
        errors.append("policy.strict_evidence must be true")
    required_metadata = policy.get("required_metadata", {})
    if not isinstance(required_metadata, dict):
        errors.append("policy.required_metadata must be an object")
        return
    metadata = packet.get("metadata")
    if not isinstance(metadata, dict):
        return
    for key, expected in required_metadata.items():
        if not isinstance(expected, str) or not expected:
            errors.append(f"policy.required_metadata.{key} must be non-empty text")
        elif metadata.get(key) != expected:
            errors.append(f"metadata.{key} must remain {expected!r} for this packet")


def _is_safe_basename(value: str) -> bool:
    return (
        bool(value)
        and Path(value).name == value
        and "/" not in value
        and "\\" not in value
    )


def _validate_input_record_schema(packet: dict[str, Any], errors: list[str]) -> None:
    record = packet.get("input_record")
    if record is None:
        return
    if not isinstance(record, dict):
        errors.append("input_record must be an object")
        return
    if record.get("kind") != "pdf":
        errors.append("input_record.kind must be 'pdf' when input_record is present")
    filename = record.get("filename")
    if (
        not isinstance(filename, str)
        or not _is_safe_basename(filename)
        or not filename.casefold().endswith(".pdf")
    ):
        errors.append("input_record.filename must be a PDF basename")
    digest = record.get("sha256")
    if not isinstance(digest, str) or not SHA256_PATTERN.fullmatch(digest):
        errors.append("input_record.sha256 must be a lowercase SHA-256 digest")
    pages = record.get("pages")
    if not isinstance(pages, int) or isinstance(pages, bool) or pages < 1:
        errors.append("input_record.pages must be a positive integer")
    for field, suffix in (
        ("extracted_text", ".txt"),
        ("manifest", ".json"),
    ):
        value = record.get(field)
        if (
            not isinstance(value, str)
            or not _is_safe_basename(value)
            or not value.casefold().endswith(suffix)
        ):
            errors.append(f"input_record.{field} must be a {suffix} basename")
    classification = record.get("classification")
    if (
        not isinstance(classification, str)
        or "not a Law Wiki legal authority" not in classification
    ):
        errors.append("input_record must state that it is not a legal authority")


def _validate_evidence(packet: dict[str, Any], errors: list[str]) -> dict[str, str]:
    matrix = packet.get("evidence_matrix")
    if not isinstance(matrix, list) or not matrix:
        errors.append("evidence_matrix must contain entries")
        return {}
    statuses: dict[str, str] = {}
    for index, item in enumerate(matrix):
        prefix = f"evidence_matrix[{index}]"
        if not isinstance(item, dict):
            errors.append(f"{prefix} must be an object")
            continue
        item_id = item.get("id")
        proposition = item.get("proposition")
        status = item.get("status")
        basis = item.get("basis")
        if not isinstance(item_id, str) or not item_id:
            errors.append(f"{prefix}.id must be non-empty text")
            continue
        if item_id in statuses:
            errors.append(f"duplicate evidence id: {item_id}")
        if not isinstance(proposition, str) or not proposition.strip():
            errors.append(f"{prefix}.proposition must be non-empty text")
        if not isinstance(status, str) or status not in ALLOWED_STATUSES:
            errors.append(f"{prefix}.status is invalid: {status}")
        if not isinstance(basis, str) or basis not in ALLOWED_BASES:
            errors.append(f"{prefix}.basis is invalid: {basis}")
        if (
            isinstance(status, str)
            and status in STATUS_BASES
            and basis not in STATUS_BASES[status]
        ):
            errors.append(f"{prefix} has incompatible status/basis: {status}/{basis}")
        statuses[item_id] = str(status)
    return statuses


def _validate_sections(
    packet: dict[str, Any],
    source_ids: set[str],
    evidence_statuses: dict[str, str],
    errors: list[str],
) -> None:
    sections = packet.get("sections")
    if not isinstance(sections, dict):
        errors.append("sections must be an object")
        return
    summary = sections.get("sumario")
    if isinstance(summary, list) and len(summary) not in {4, 5}:
        errors.append("sections.sumario must contain four or five points")
    legal_source_citations = 0
    for section_key in BLOCK_MARKERS.values():
        content = sections.get(section_key)
        if not isinstance(content, list) or not content:
            errors.append(f"sections.{section_key} must contain paragraphs")
            continue
        for index, spec in enumerate(content):
            text = _paragraph_text(spec)
            prefix = f"sections.{section_key}[{index}]"
            if not text.strip():
                errors.append(f"{prefix} has no text")
            if not isinstance(spec, dict):
                if section_key in {"factos_provados", "factos_nao_provados"}:
                    errors.append(f"{prefix} must cite evidence_ids")
                continue
            cited = spec.get("source_ids", [])
            if not isinstance(cited, list):
                errors.append(f"{prefix}.source_ids must be a list")
                cited = []
            else:
                invalid_citations = [item for item in cited if not isinstance(item, str)]
                if invalid_citations:
                    errors.append(f"{prefix}.source_ids must contain text ids")
                unknown = {item for item in cited if isinstance(item, str)} - source_ids
                if unknown:
                    errors.append(f"{prefix} cites unknown sources: {sorted(unknown)}")
            if section_key in {"motivacao", "fundamentacao_direito"}:
                legal_source_citations += len(cited)
            evidence_ids = spec.get("evidence_ids", [])
            if not isinstance(evidence_ids, list):
                errors.append(f"{prefix}.evidence_ids must be a list")
                evidence_ids = []
            invalid_evidence = [item for item in evidence_ids if not isinstance(item, str)]
            if invalid_evidence:
                errors.append(f"{prefix}.evidence_ids must contain text ids")
            unknown_evidence = {
                item for item in evidence_ids if isinstance(item, str)
            } - set(evidence_statuses)
            if unknown_evidence:
                errors.append(f"{prefix} cites unknown evidence: {sorted(unknown_evidence)}")
            if section_key == "factos_provados":
                if not evidence_ids:
                    errors.append(f"{prefix} needs evidence_ids")
                for evidence_id in evidence_ids:
                    if evidence_statuses.get(evidence_id) not in {"proved", "procedural"}:
                        errors.append(
                            f"{prefix} relies on non-proved evidence {evidence_id}"
                        )
            if section_key == "factos_nao_provados":
                if not evidence_ids:
                    errors.append(f"{prefix} needs evidence_ids")
                for evidence_id in evidence_ids:
                    if evidence_statuses.get(evidence_id) != "not_proved":
                        errors.append(
                            f"{prefix} lists proved evidence {evidence_id} as unproved"
                        )
    if legal_source_citations == 0:
        errors.append("motivacao/fundamentacao_direito must cite local authorities")


def _decimal_value(value: Any, label: str, errors: list[str]) -> Decimal:
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError):
        errors.append(f"{label} must be a decimal amount")
        return Decimal("0")


def _validate_amounts(packet: dict[str, Any], errors: list[str]) -> None:
    amounts = packet.get("amounts")
    if not isinstance(amounts, list) or not amounts:
        errors.append("amounts must contain itemized claims")
        return
    awarded_total = Decimal("0")
    claimed_total = Decimal("0")
    for index, amount in enumerate(amounts):
        prefix = f"amounts[{index}]"
        if not isinstance(amount, dict):
            errors.append(f"{prefix} must be an object")
            continue
        if not isinstance(amount.get("item"), str) or not amount["item"].strip():
            errors.append(f"{prefix}.item must be non-empty text")
        claimed = _decimal_value(amount.get("claimed"), f"{prefix}.claimed", errors)
        awarded = _decimal_value(amount.get("awarded"), f"{prefix}.awarded", errors)
        if claimed < 0 or awarded < 0:
            errors.append(f"{prefix} cannot contain negative amounts")
        if awarded > claimed:
            errors.append(f"{prefix}.awarded exceeds the pleaded amount")
        if not isinstance(amount.get("reason"), str) or not amount["reason"].strip():
            errors.append(f"{prefix}.reason must be non-empty text")
        claimed_total += claimed
        awarded_total += awarded
    expected = _decimal_value(packet.get("decision_total"), "decision_total", errors)
    if awarded_total.quantize(Decimal("0.01")) != expected.quantize(Decimal("0.01")):
        errors.append(
            f"award arithmetic mismatch: items={awarded_total}, total={expected}"
        )
    metadata = packet.get("metadata")
    if isinstance(metadata, dict):
        case_value = _decimal_value(
            metadata.get("case_value"), "metadata.case_value", errors
        )
        if claimed_total.quantize(Decimal("0.01")) != case_value.quantize(Decimal("0.01")):
            errors.append(
                f"case value mismatch: claims={claimed_total}, "
                f"metadata.case_value={case_value}"
            )


def validate_packet(packet: dict[str, Any], *, catalog: dict[str, Any]) -> list[str]:
    """Return human-readable packet errors. An empty list means the packet is valid.

    Sources are checked against ``catalog`` (document id exists, path matches
    ``path`` or ``rel_source``). Raw files are not opened. Paragraph keys
    ``id``, ``confidence``, and ``note`` are ignored.
    """
    if not isinstance(packet, dict):
        return ["packet must be an object"]
    errors: list[str] = []
    _validate_metadata(packet, errors)
    _validate_policy(packet, errors)
    _validate_input_record_schema(packet, errors)
    source_errors, sources = _collect_sources(packet, catalog)
    errors.extend(source_errors)
    evidence_statuses = _validate_evidence(packet, errors)
    _validate_sections(packet, set(sources), evidence_statuses, errors)
    _validate_amounts(packet, errors)
    all_text = "\n".join(_iter_packet_text(packet))
    if re.search(r"https?://|file://", all_text, flags=re.IGNORECASE):
        errors.append("packet contains an external URL")
    if re.search(r"\[\[[A-Z_]+\]\]|\(\s*indicar\b", all_text, flags=re.IGNORECASE):
        errors.append("packet contains unresolved template instructions")
    return errors


def _expected_public_urls(packet: dict[str, Any]) -> set[str]:
    urls: set[str] = set()
    sources = packet.get("sources")
    if not isinstance(sources, list):
        return urls
    for source in sources:
        if not isinstance(source, dict):
            continue
        document_id = source.get("document_id")
        if not isinstance(document_id, str):
            continue
        try:
            urls.add(public_document_url(document_id))
        except ValueError:
            continue
    return urls


def _docx_text(path: Path) -> str:
    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for section in document.sections:
        parts.extend(paragraph.text for paragraph in section.header.paragraphs)
        parts.extend(paragraph.text for paragraph in section.footer.paragraphs)
    return "\n".join(parts)


def _validate_rendered_text(text: str, label: str, errors: list[str]) -> None:
    if WARNING not in text:
        errors.append(f"{label} is missing the unsigned-draft warning")
    for heading in REQUIRED_HEADINGS:
        if heading not in text:
            errors.append(f"{label} is missing heading: {heading}")
    if re.search(r"\[\[[A-Z_]+\]\]|\(\s*indicar\b", text, flags=re.IGNORECASE):
        errors.append(f"{label} contains unresolved template markers")
    if "Sem identidade ou assinatura aposta" not in text:
        errors.append(f"{label} is missing the no-signature statement")


def _relationship_part_name(part_name: str) -> str:
    path = Path(part_name)
    return str(path.parent / "_rels" / f"{path.name}.rels").replace("\\", "/")


def _docx_hyperlinks(path: Path) -> set[str]:
    targets: set[str] = set()
    with zipfile.ZipFile(path) as archive:
        part_names = [
            name
            for name in archive.namelist()
            if name == "word/document.xml"
            or re.fullmatch(r"word/(?:header|footer)\d+\.xml", name)
        ]
        for part_name in part_names:
            root = ElementTree.fromstring(archive.read(part_name))
            relationship_ids = {
                node.get(f"{{{OFFICE_RELATIONSHIPS_NAMESPACE}}}id")
                for node in root.iter(f"{{{WORD_NAMESPACE}}}hyperlink")
            }
            relationship_ids.discard(None)
            rels_name = _relationship_part_name(part_name)
            if rels_name not in archive.namelist():
                continue
            rels = ElementTree.fromstring(archive.read(rels_name))
            for relation in rels:
                if relation.get("Id") not in relationship_ids:
                    continue
                if relation.get("TargetMode") != "External":
                    continue
                target = relation.get("Target")
                if target:
                    targets.add(target)
    return targets


def _pdf_hyperlinks(reader: PdfReader) -> set[str]:
    targets: set[str] = set()
    for page in reader.pages:
        annotations = page.get("/Annots") or []
        for annotation_ref in annotations:
            annotation = annotation_ref.get_object()
            action = annotation.get("/A")
            if action is None:
                continue
            action = action.get_object() if hasattr(action, "get_object") else action
            uri = action.get("/URI") if hasattr(action, "get") else None
            if uri:
                targets.add(str(uri))
    return targets


def _validate_hyperlinks(
    targets: set[str],
    expected_urls: set[str],
    label: str,
    errors: list[str],
) -> None:
    missing = expected_urls - targets
    if missing:
        errors.append(f"{label} is missing public source hyperlinks: {sorted(missing)}")
    forbidden = {
        target
        for target in targets
        if target.startswith(("file://", "http://", "https://"))
        and not target.startswith(f"{PUBLIC_BASE_URL}/documents/")
    }
    if forbidden:
        errors.append(f"{label} contains non-Law-Wiki hyperlinks: {sorted(forbidden)}")


def validate_docx(path: Path, packet: dict[str, Any]) -> list[str]:
    """Check a rendered DOCX for the warning, headings, and public source links."""
    path = Path(path)
    errors: list[str] = []
    expected_urls = _expected_public_urls(packet)
    if not path.is_file() or path.stat().st_size == 0:
        errors.append(f"DOCX output is missing or empty: {path}")
        return errors
    _validate_rendered_text(_docx_text(path), "DOCX", errors)
    try:
        _validate_hyperlinks(_docx_hyperlinks(path), expected_urls, "DOCX", errors)
    except (OSError, zipfile.BadZipFile, ElementTree.ParseError) as exc:
        errors.append(f"DOCX hyperlinks cannot be inspected: {exc}")
    return errors


def validate_pdf(path: Path, packet: dict[str, Any]) -> list[str]:
    """Check a rendered PDF for the warning, headings, and public source links."""
    path = Path(path)
    errors: list[str] = []
    expected_urls = _expected_public_urls(packet)
    if not path.is_file() or path.stat().st_size == 0:
        errors.append(f"PDF output is missing or empty: {path}")
        return errors
    try:
        reader = PdfReader(str(path))
        if not reader.pages:
            errors.append("PDF has no pages")
            return errors
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        _validate_rendered_text(text, "PDF", errors)
        _validate_hyperlinks(_pdf_hyperlinks(reader), expected_urls, "PDF", errors)
    except Exception as exc:
        errors.append(f"PDF cannot be parsed: {exc}")
    return errors


def _manifest_expectations(packet: dict[str, Any], errors: list[str]) -> dict[str, dict[str, Any]]:
    sources = packet.get("sources")
    if not isinstance(sources, list):
        errors.append("sources must contain at least one local authority")
        return {}
    expected: dict[str, dict[str, Any]] = {}
    for index, source in enumerate(sources):
        prefix = f"sources[{index}]"
        if not isinstance(source, dict):
            errors.append(f"{prefix} must be an object")
            continue
        source_id = source.get("id")
        if not isinstance(source_id, str) or not source_id.strip():
            errors.append(f"{prefix}.id must be non-empty text")
            continue
        if source_id in expected:
            errors.append(f"duplicate source id: {source_id}")
            continue
        try:
            path_value = source.get("path")
            document_id = source.get("document_id")
            if not isinstance(path_value, str) or not isinstance(document_id, str):
                raise ValueError(
                    "Every source needs non-empty id, title, path, and document_id fields"
                )
            normalized = repository_path(path_value)
            public_url = public_document_url(document_id)
        except ValueError as exc:
            errors.append(str(exc))
            continue
        expected[source_id] = {
            "title": source.get("title"),
            "document_id": document_id,
            "repository_path": normalized,
            "pinpoints": source.get("pinpoints"),
            "public_url": public_url,
        }
    return expected


def validate_source_manifest(manifest: dict[str, Any] | Path, packet: dict[str, Any]) -> list[str]:
    """Check a source manifest dict, or a JSON file, against the packet.

    Does not open raw corpus files.
    """
    errors: list[str] = []
    if isinstance(manifest, Path) or (
        isinstance(manifest, str) and not isinstance(manifest, dict)
    ):
        try:
            payload = _load_json(Path(manifest))
        except ValueError as exc:
            return [str(exc)]
    elif isinstance(manifest, dict):
        payload = manifest
    else:
        return ["source manifest must be an object or a path"]
    if payload.get("warning") != WARNING:
        errors.append("source manifest has the wrong unsigned-draft warning")
    if payload.get("browse_url") != PUBLIC_FOLDERS_URL:
        errors.append("source manifest has the wrong folders browse URL")
    text = json.dumps(payload, ensure_ascii=False)
    if "file://" in text.lower():
        errors.append("source manifest exposes a file:// URL")
    manifest_sources = payload.get("sources")
    if not isinstance(manifest_sources, list):
        errors.append("source manifest sources must be a list")
        return errors
    by_id = {
        item.get("id"): item
        for item in manifest_sources
        if isinstance(item, dict) and isinstance(item.get("id"), str)
    }
    expected = _manifest_expectations(packet, errors)
    if set(by_id) != set(expected):
        errors.append("source manifest ids do not match packet sources")
    for source_id, source in expected.items():
        item = by_id.get(source_id)
        if not isinstance(item, dict):
            continue
        for key, value in source.items():
            if item.get(key) != value:
                errors.append(f"manifest source {source_id} has wrong {key}")
    return errors


def validate_input_manifest(path: Path | None, packet: dict[str, Any]) -> list[str]:
    """Check an extracted PDF input manifest against ``packet['input_record']``."""
    errors: list[str] = []
    record = packet.get("input_record")
    if record is None:
        if path is not None:
            errors.append("input manifest was supplied but packet has no input_record")
        return errors
    if not isinstance(record, dict) or record.get("kind") != "pdf":
        return errors
    if path is None:
        errors.append("PDF input packet requires --input-manifest")
        return errors
    path = Path(path)
    if path.name != record.get("manifest"):
        errors.append("input manifest filename does not match packet input_record")
    try:
        manifest = _load_json(path)
    except ValueError as exc:
        errors.append(str(exc))
        return errors
    expected_fields = {
        "kind": "case-record-input",
        "filename": record.get("filename"),
        "sha256": record.get("sha256"),
        "pages": record.get("pages"),
        "classification": record.get("classification"),
    }
    for key, expected in expected_fields.items():
        if manifest.get(key) != expected:
            errors.append(f"input manifest has wrong {key}")
    text_name = record.get("extracted_text")
    if isinstance(text_name, str) and _is_safe_basename(text_name):
        text_path = path.parent / text_name
        if not text_path.is_file() or text_path.stat().st_size == 0:
            errors.append(f"extracted PDF record is missing or empty: {text_path}")
    return errors
