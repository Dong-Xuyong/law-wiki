"""Render a grounded Law Wiki decision packet to DOCX.

Legal authorities are resolved from a catalog mapping. Raw corpus files are
not required on disk. PDF export is optional and imports Word COM only when
called.
"""

from __future__ import annotations

import json
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt


WARNING = "MINUTA NÃO ASSINADA — AUXÍLIO À INVESTIGAÇÃO JURÍDICA"
PUBLIC_BASE_URL = "https://dong-xuyong.github.io/law-wiki"
PUBLIC_FOLDERS_URL = f"{PUBLIC_BASE_URL}/folders/"
DOCUMENT_ID_PATTERN = re.compile(r"^law_[0-9a-f]{16}$")
BLOCK_MARKERS = {
    "SUMARIO": "sumario",
    "RELATORIO": "relatorio",
    "OBJETO_LITIGIO": "objeto_litigio",
    "SANEADOR": "saneador",
    "FACTOS_PROVADOS": "factos_provados",
    "FACTOS_NAO_PROVADOS": "factos_nao_provados",
    "MOTIVACAO": "motivacao",
    "FUNDAMENTACAO_DIREITO": "fundamentacao_direito",
    "DECISAO": "decisao",
}
SCALAR_MARKERS = {
    "PROCESSO_NUMERO": ("metadata", "process_number"),
    "RECLAMANTE": ("metadata", "claimant"),
    "RECLAMADA": ("metadata", "respondent"),
    "VALOR_PROCESSO": ("metadata", "case_value"),
    "LOCAL_DATA": ("metadata", "place_and_date"),
}
ALIGNMENTS = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
IGNORED_PARAGRAPH_KEYS = {"id", "confidence", "note"}
CATALOG_FILES = ("catalog.jsonl", "open-sources-catalog.jsonl")


def repository_path(source_path: str) -> str:
    value = source_path.replace("\\", "/").removeprefix("./")
    if re.match(r"^[a-z]+://", value, flags=re.IGNORECASE):
        raise ValueError(f"External source is forbidden: {source_path}")
    relative = Path(value)
    if relative.is_absolute() or ".." in relative.parts:
        raise ValueError(f"Source path must be repository-relative: {source_path}")
    if not relative.parts or relative.parts[0] not in {"raw", "wiki"}:
        raise ValueError(f"Source must be inside raw/ or wiki/: {source_path}")
    if relative.parts[0] == "raw" and (
        len(relative.parts) < 2 or relative.parts[1] != "documents"
    ):
        raise ValueError(f"Raw authority must be inside raw/documents/: {source_path}")
    return relative.as_posix()


def public_document_url(document_id: str) -> str:
    if not DOCUMENT_ID_PATTERN.fullmatch(document_id):
        raise ValueError(f"Invalid Law Wiki document id: {document_id}")
    return f"{PUBLIC_BASE_URL}/documents/{document_id}/"


def _normalize_catalog_path(value: str) -> str:
    return value.replace("\\", "/").removeprefix("./")


def _row_paths(row: dict[str, Any]) -> dict[str, str]:
    rel_raw = row.get("rel_raw")
    if not isinstance(rel_raw, str) or not rel_raw.strip():
        rel_raw = row.get("path") if isinstance(row.get("path"), str) else ""
    rel_source = row.get("rel_source") if isinstance(row.get("rel_source"), str) else ""
    return {
        "rel_raw": _normalize_catalog_path(str(rel_raw or "")),
        "rel_source": _normalize_catalog_path(rel_source),
    }


def load_local_catalog(root: Path) -> dict[str, dict[str, Any]]:
    """Read ``data/catalog.jsonl`` and, when present, ``data/open-sources-catalog.jsonl``.

    Returns ``{document_id: {"path": rel_raw, "title": title, "rel_source": rel_source}}``.
    ``rel_source`` is ``None`` when the row has none. Raw files are not opened.
    """
    root = Path(root)
    data_dir = root / "data"
    seen_catalog = False
    by_id: dict[str, dict[str, str]] = {}
    by_path: dict[str, str] = {}
    for name in CATALOG_FILES:
        catalog_path = data_dir / name
        if not catalog_path.is_file():
            continue
        seen_catalog = True
        for line_number, line in enumerate(
            catalog_path.read_text(encoding="utf-8").splitlines(), start=1
        ):
            if not line.strip():
                continue
            try:
                row = json.loads(line)
            except json.JSONDecodeError as exc:
                raise ValueError(
                    f"Invalid catalog JSON at {catalog_path}:{line_number}: {exc}"
                ) from exc
            if not isinstance(row, dict):
                raise ValueError(
                    f"Catalog row must be an object at {catalog_path}:{line_number}"
                )
            document_id = str(row.get("id") or "")
            if not DOCUMENT_ID_PATTERN.fullmatch(document_id):
                continue
            paths = _row_paths(row)
            previous = by_id.get(document_id)
            if previous and any(
                previous[field] != paths[field]
                for field in ("rel_raw", "rel_source")
                if previous[field] and paths[field]
            ):
                raise ValueError(f"Conflicting catalog rows for {document_id}")
            by_id.setdefault(document_id, paths | {"title": str(row.get("title") or "")})
            for field in ("rel_raw", "rel_source"):
                value = paths[field]
                if not value:
                    continue
                existing = by_path.get(value)
                if existing and existing != document_id:
                    raise ValueError(
                        f"Catalog path maps to multiple documents: {value}"
                    )
                by_path[value] = document_id
    if not seen_catalog:
        raise ValueError("No Law Wiki catalog found under data/")
    return {
        document_id: {
            "path": fields["rel_raw"],
            "title": fields["title"],
            "rel_source": fields["rel_source"] or None,
        }
        for document_id, fields in by_id.items()
    }


def _entry_paths(entry: dict[str, Any]) -> list[str]:
    paths: list[str] = []
    raw = entry.get("path")
    if not isinstance(raw, str) or not raw.strip():
        raw = entry.get("rel_raw") if isinstance(entry.get("rel_raw"), str) else ""
    if isinstance(raw, str) and raw.strip():
        paths.append(_normalize_catalog_path(raw))
    rel_source = entry.get("rel_source")
    if isinstance(rel_source, str) and rel_source.strip():
        paths.append(_normalize_catalog_path(rel_source))
    return [path for path in paths if path]


def _path_index(catalog: dict[str, Any]) -> dict[str, str]:
    index: dict[str, str] = {}
    for document_id, entry in catalog.items():
        if not isinstance(document_id, str) or not isinstance(entry, dict):
            continue
        for path in _entry_paths(entry):
            index.setdefault(path, document_id)
    return index


def _enrich_source(
    source: dict[str, Any],
    catalog: dict[str, Any],
    path_index: dict[str, str],
) -> dict[str, Any]:
    source_id = source.get("id")
    title = source.get("title")
    path_value = source.get("path")
    document_id = source.get("document_id")
    if not all(
        isinstance(item, str) and item.strip()
        for item in (source_id, title, path_value, document_id)
    ):
        raise ValueError(
            "Every source needs non-empty id, title, path, and document_id fields"
        )
    assert isinstance(source_id, str)
    assert isinstance(title, str)
    assert isinstance(path_value, str)
    assert isinstance(document_id, str)
    normalized_path = repository_path(path_value)
    if not DOCUMENT_ID_PATTERN.fullmatch(document_id):
        raise ValueError(f"Source {source_id} has invalid document_id: {document_id}")
    owner = path_index.get(normalized_path)
    if owner is None:
        raise ValueError(
            f"Source path is absent from the Law Wiki catalog: {normalized_path}"
        )
    if owner != document_id:
        raise ValueError(
            f"Source {source_id} path resolves to {owner}, not {document_id}"
        )
    if document_id not in catalog:
        raise ValueError(
            f"Source document id is absent from the catalog: {document_id}"
        )
    pinpoints = source.get("pinpoints")
    if (
        not isinstance(pinpoints, list)
        or not pinpoints
        or not all(isinstance(item, str) and item.strip() for item in pinpoints)
    ):
        raise ValueError(f"Source {source_id} needs non-empty textual pinpoints")
    return {
        "id": source_id,
        "title": title,
        "path": path_value,
        "document_id": document_id,
        "pinpoints": list(pinpoints),
        "repository_path": normalized_path,
        "public_url": public_document_url(document_id),
    }


def _collect_sources(
    packet: dict[str, Any], catalog: dict[str, Any]
) -> tuple[list[str], dict[str, dict[str, Any]]]:
    errors: list[str] = []
    sources = packet.get("sources")
    if not isinstance(sources, list) or not sources:
        errors.append("sources must contain at least one local authority")
        return errors, {}
    if not isinstance(catalog, dict):
        errors.append("catalog must be an object")
        return errors, {}
    path_index = _path_index(catalog)
    enriched: dict[str, dict[str, Any]] = {}
    for index, source in enumerate(sources):
        prefix = f"sources[{index}]"
        if not isinstance(source, dict):
            errors.append(f"{prefix} must be an object")
            continue
        source_id = source.get("id")
        if not isinstance(source_id, str) or not source_id.strip():
            errors.append(f"{prefix}.id must be non-empty text")
            continue
        if source_id in enriched:
            errors.append(f"duplicate source id: {source_id}")
            continue
        try:
            resolved = _enrich_source(source, catalog, path_index)
        except ValueError as exc:
            errors.append(str(exc))
            continue
        enriched[source_id] = resolved
    return errors, enriched


def _source_manifest(
    packet: dict[str, Any], sources: dict[str, dict[str, Any]]
) -> dict[str, Any]:
    metadata = packet.get("metadata", {})
    if not isinstance(metadata, dict):
        metadata = {}
    return {
        "document": metadata.get("document_title", "Minuta de decisão arbitral"),
        "draft_date": metadata.get("draft_date"),
        "database_policy": (
            "Legal propositions and jurisprudence are limited to the local Law Wiki."
        ),
        "warning": WARNING,
        "browse_url": PUBLIC_FOLDERS_URL,
        "sources": [
            {
                "id": source["id"],
                "title": source["title"],
                "document_id": source["document_id"],
                "repository_path": source["repository_path"],
                "pinpoints": source["pinpoints"],
                "public_url": source["public_url"],
            }
            for source in sources.values()
        ],
    }


def _iter_story_paragraphs(document: Document) -> Iterable[Any]:
    yield from document.paragraphs
    for section in document.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def _get_nested(packet: dict[str, Any], keys: tuple[str, ...]) -> str:
    value: Any = packet
    for key in keys:
        if not isinstance(value, dict) or key not in value:
            raise ValueError(f"Missing packet field: {'.'.join(keys)}")
        value = value[key]
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"Packet field must be non-empty text: {'.'.join(keys)}")
    return value.strip()


def _format_currency_pt(value: str) -> str:
    try:
        amount = Decimal(value).quantize(Decimal("0.01"))
    except InvalidOperation:
        return value
    sign = "-" if amount < 0 else ""
    whole, fraction = f"{abs(amount):.2f}".split(".")
    grouped = f"{int(whole):,}".replace(",", " ")
    return f"{sign}€ {grouped},{fraction}"


def _replace_scalar_markers(document: Document, packet: dict[str, Any]) -> None:
    replacements: dict[str, str] = {}
    for marker, keys in SCALAR_MARKERS.items():
        value = _get_nested(packet, keys)
        if marker == "VALOR_PROCESSO":
            value = _format_currency_pt(value)
        replacements[f"[[{marker}]]"] = value
    found = {marker: False for marker in replacements}
    for paragraph in _iter_story_paragraphs(document):
        original = paragraph.text
        revised = original
        for marker, value in replacements.items():
            if marker in revised:
                revised = revised.replace(marker, value)
                found[marker] = True
        if revised != original:
            if paragraph.runs:
                paragraph.runs[0].text = revised
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.add_run(revised)
    missing = [marker for marker, seen in found.items() if not seen]
    if missing:
        raise ValueError(f"Template is missing scalar markers: {', '.join(missing)}")


def _add_hyperlink(paragraph: Any, label: str, target: str) -> None:
    relationship_id = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend((color, underline))
    run.append(properties)
    text = OxmlElement("w:t")
    text.set(qn("xml:space"), "preserve")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _apply_paragraph_spec(
    paragraph: Any,
    spec: str | dict[str, Any],
    sources: dict[str, dict[str, Any]],
) -> None:
    if isinstance(spec, str):
        text = spec
        options: dict[str, Any] = {}
    elif isinstance(spec, dict):
        options = {
            key: value
            for key, value in spec.items()
            if key not in IGNORED_PARAGRAPH_KEYS
        }
        text = options.get("text", "")
    else:
        raise ValueError("Section paragraphs must be strings or objects")
    if not isinstance(text, str) or not text.strip():
        raise ValueError("Section paragraphs cannot be empty")

    style = options.get("style")
    if style:
        paragraph.style = style
    alignment = options.get("align", "justify")
    if alignment not in ALIGNMENTS:
        raise ValueError(f"Unsupported paragraph alignment: {alignment}")
    paragraph.alignment = ALIGNMENTS[alignment]
    if "indent_cm" in options:
        paragraph.paragraph_format.left_indent = Cm(float(options["indent_cm"]))
    if "first_line_cm" in options:
        paragraph.paragraph_format.first_line_indent = Cm(
            float(options["first_line_cm"])
        )
    if "space_after_pt" in options:
        paragraph.paragraph_format.space_after = Pt(float(options["space_after_pt"]))
    if options.get("keep_with_next"):
        paragraph.paragraph_format.keep_with_next = True

    run = paragraph.add_run(text.strip())
    run.bold = bool(options.get("bold", False))
    run.italic = bool(options.get("italic", False))

    source_ids = options.get("source_ids", [])
    if not isinstance(source_ids, list):
        raise ValueError("source_ids must be a list")
    for source_id in source_ids:
        if source_id not in sources:
            raise ValueError(f"Paragraph cites unknown source id: {source_id}")
        source = sources[source_id]
        suffix = f", {', '.join(source['pinpoints'])}"
        paragraph.add_run(" ")
        _add_hyperlink(
            paragraph,
            f"[Fonte Law Wiki: {source['title']}{suffix}]",
            str(source["public_url"]),
        )


def _replace_block_markers(
    document: Document,
    packet: dict[str, Any],
    sources: dict[str, dict[str, Any]],
) -> None:
    sections = packet.get("sections")
    if not isinstance(sections, dict):
        raise ValueError("Decision packet must include a sections object")

    paragraphs_by_text = {
        paragraph.text.strip(): paragraph for paragraph in document.paragraphs
    }
    for marker, section_key in BLOCK_MARKERS.items():
        marker_text = f"[[{marker}]]"
        marker_paragraph = paragraphs_by_text.get(marker_text)
        if marker_paragraph is None:
            raise ValueError(f"Template is missing block marker: {marker_text}")
        content = sections.get(section_key)
        if not isinstance(content, list) or not content:
            raise ValueError(f"Section must contain paragraph entries: {section_key}")
        for spec in content:
            paragraph = marker_paragraph.insert_paragraph_before()
            _apply_paragraph_spec(paragraph, spec, sources)
        marker_paragraph._element.getparent().remove(marker_paragraph._element)


def render_docx(
    packet: dict[str, Any],
    *,
    template: Path,
    out_docx: Path,
    catalog: dict[str, Any],
) -> dict[str, Any]:
    """Render ``packet`` to ``out_docx`` and return the source manifest.

    The manifest lists ``repository_path``, ``document_id``, ``pinpoints``, and
    ``public_url`` (``https://dong-xuyong.github.io/law-wiki/documents/{id}/``).
    It contains no ``file://`` links. Paragraph keys ``id``, ``confidence``,
    and ``note`` are ignored.
    """
    if not isinstance(packet, dict):
        raise ValueError("Decision packet must be a JSON object")
    errors, sources = _collect_sources(packet, catalog)
    if errors:
        raise ValueError("; ".join(errors))
    template = Path(template)
    out_docx = Path(out_docx)
    if not template.is_file():
        raise ValueError(f"Template does not exist: {template}")
    document = Document(str(template))
    _replace_scalar_markers(document, packet)
    _replace_block_markers(document, packet, sources)
    out_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out_docx))
    return _source_manifest(packet, sources)


def _update_word_fields(document: Any) -> None:
    document.Fields.Update()
    for section in document.Sections:
        for header in section.Headers:
            if header.Exists:
                header.Range.Fields.Update()
        for footer in section.Footers:
            if footer.Exists:
                footer.Range.Fields.Update()


def export_pdf(docx: Path, pdf: Path) -> None:
    """Export ``docx`` to ``pdf`` via Microsoft Word. Windows-only.

    ``win32com`` is imported inside this function, never at module import.
    """
    docx = Path(docx)
    pdf = Path(pdf)
    try:
        import win32com.client  # type: ignore[import-not-found]
    except ImportError as exc:
        raise RuntimeError(
            "PDF export requires pywin32 and Microsoft Word on Windows"
        ) from exc

    word = None
    opened = None
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        opened = word.Documents.Open(
            str(docx.resolve()),
            ConfirmConversions=False,
            ReadOnly=False,
            AddToRecentFiles=False,
        )
        _update_word_fields(opened)
        opened.Save()
        pdf.parent.mkdir(parents=True, exist_ok=True)
        opened.ExportAsFixedFormat(
            OutputFileName=str(pdf.resolve()),
            ExportFormat=17,
            OpenAfterExport=False,
            OptimizeFor=0,
            Range=0,
            Item=0,
            IncludeDocProps=True,
            KeepIRM=True,
            CreateBookmarks=1,
            DocStructureTags=True,
            BitmapMissingFonts=True,
            UseISO19005_1=False,
        )
    finally:
        if opened is not None:
            opened.Close(SaveChanges=False)
        if word is not None:
            word.Quit()
