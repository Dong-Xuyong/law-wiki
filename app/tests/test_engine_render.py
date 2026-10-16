"""Engine render, validate, and extract tests. No network and no Microsoft Word."""

from __future__ import annotations

import ast
import copy
import hashlib
import json
import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.engine.extract import extract_pdf_bytes
from app.engine.render import WARNING, load_local_catalog, render_docx
from app.engine.validate import validate_docx, validate_packet, validate_source_manifest


APP = Path(__file__).resolve().parents[1]
ROOT = APP.parent
TEMPLATE = APP / "engine" / "templates" / "decisao-arbitral-wiki.docx"
EXAMPLE_PACKET = (
    ROOT / ".cursor" / "skills" / "law-wiki" / "examples" / "wizz-w4-5520-packet.json"
)
EXAMPLE_PDF = (
    ROOT / ".cursor" / "skills" / "law-wiki" / "examples" / "wizz-w4-5520-input.pdf"
)

DOCUMENT_ID = "law_0123456789abcdef"
RAW_PATH = f"raw/documents/test/2026/{DOCUMENT_ID}.md"
WIKI_PATH = f"wiki/sources/test--{DOCUMENT_ID}.md"


def _document_text(path: Path) -> str:
    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for section in document.sections:
        parts.extend(paragraph.text for paragraph in section.header.paragraphs)
        parts.extend(paragraph.text for paragraph in section.footer.paragraphs)
    return "\n".join(parts)


def _text_pdf(pages: list[str]) -> bytes:
    writer = PdfWriter()
    for text in pages:
        page = writer.add_blank_page(width=300, height=144)
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = DecodedStreamObject()
        stream.set_data(f"BT /F1 12 Tf 36 80 Td ({escaped}) Tj ET".encode("latin-1"))
        font = DictionaryObject()
        font[NameObject("/Type")] = NameObject("/Font")
        font[NameObject("/Subtype")] = NameObject("/Type1")
        font[NameObject("/BaseFont")] = NameObject("/Helvetica")
        fonts = DictionaryObject()
        fonts[NameObject("/F1")] = font
        resources = DictionaryObject()
        resources[NameObject("/Font")] = fonts
        page[NameObject("/Resources")] = resources
        page[NameObject("/Contents")] = stream
    buffer = BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


class EngineRenderTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.catalog = {
            DOCUMENT_ID: {
                "path": RAW_PATH,
                "title": "Test authority",
                "rel_source": WIKI_PATH,
            }
        }
        self.packet = self.make_packet()

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def make_packet(self) -> dict:
        source = {
            "id": "authority",
            "title": "Test authority",
            "path": RAW_PATH,
            "document_id": DOCUMENT_ID,
            "pinpoints": ["linhas 1–3"],
        }
        cited = {"text": "Texto jurídico local.", "source_ids": ["authority"]}
        return {
            "metadata": {
                "process_number": "42/2026",
                "claimant": "Parte A",
                "respondent": "Parte B",
                "case_value": "1.00",
                "place_and_date": "Porto, 16 de setembro de 2026",
                "draft_date": "2026-09-16",
                "document_title": "Minuta de teste",
            },
            "policy": {
                "strict_evidence": True,
                "required_metadata": {"process_number": "42/2026"},
            },
            "sources": [source],
            "evidence_matrix": [
                {
                    "id": "p1",
                    "proposition": "Facto comum.",
                    "status": "proved",
                    "basis": "common",
                },
                {
                    "id": "u1",
                    "proposition": "Facto unilateral.",
                    "status": "not_proved",
                    "basis": "unilateral_assertion",
                },
            ],
            "sections": {
                "sumario": [
                    cited,
                    {"text": "2. Segundo ponto."},
                    {"text": "3. Terceiro ponto."},
                    {"text": "4. Quarto ponto."},
                ],
                "relatorio": [{"text": "Relatório."}],
                "objeto_litigio": [{"text": "Objeto."}],
                "saneador": [{"text": "Saneador."}],
                "factos_provados": [{"text": "a) Facto.", "evidence_ids": ["p1"]}],
                "factos_nao_provados": [
                    {"text": "aa) Facto.", "evidence_ids": ["u1"]}
                ],
                "motivacao": [cited],
                "fundamentacao_direito": [cited],
                "decisao": [{"text": "Julga-se parcialmente procedente."}],
            },
            "amounts": [
                {
                    "item": "Pedido",
                    "claimed": "1.00",
                    "awarded": "1.00",
                    "reason": "Provado.",
                }
            ],
            "decision_total": "1.00",
        }

    def render(self, packet: dict | None = None, name: str = "decision.docx") -> dict:
        return render_docx(
            self.packet if packet is None else packet,
            template=TEMPLATE,
            out_docx=self.root / name,
            catalog=self.catalog,
        )

    def test_public_url_is_catalog_backed(self) -> None:
        manifest = self.render()
        self.assertEqual(
            manifest["sources"][0]["public_url"],
            f"https://dong-xuyong.github.io/law-wiki/documents/{DOCUMENT_ID}/",
        )
        bad = copy.deepcopy(self.packet)
        bad["sources"][0]["document_id"] = "law_ffffffffffffffff"
        with self.assertRaisesRegex(ValueError, "not law_ffffffffffffffff"):
            self.render(bad, "bad.docx")

    def test_render_uses_public_links_and_manifest_has_no_file_uri(self) -> None:
        docx = self.root / "decision.docx"
        manifest = render_docx(
            self.packet,
            template=TEMPLATE,
            out_docx=docx,
            catalog=self.catalog,
        )
        expected = f"https://dong-xuyong.github.io/law-wiki/documents/{DOCUMENT_ID}/"
        self.assertEqual(manifest["browse_url"], "https://dong-xuyong.github.io/law-wiki/folders/")
        self.assertEqual(manifest["sources"][0]["public_url"], expected)
        self.assertEqual(manifest["sources"][0]["repository_path"], RAW_PATH)
        self.assertEqual(manifest["sources"][0]["document_id"], DOCUMENT_ID)
        self.assertEqual(manifest["sources"][0]["pinpoints"], ["linhas 1–3"])
        self.assertNotIn("file://", json.dumps(manifest).lower())
        self.assertEqual(validate_packet(self.packet, catalog=self.catalog), [])
        self.assertEqual(validate_source_manifest(manifest, self.packet), [])
        self.assertEqual(validate_docx(docx, self.packet), [])
        text = _document_text(docx)
        self.assertIn("[Fonte Law Wiki: Test authority, linhas 1–3]", text)
        self.assertIn(WARNING, text)

    def test_catalog_only_mode_without_raw_files(self) -> None:
        self.assertFalse((self.root / RAW_PATH).exists())
        self.assertFalse((ROOT / RAW_PATH).exists())
        errors = validate_packet(self.packet, catalog=self.catalog)
        self.assertEqual(errors, [])
        manifest = self.render(name="catalog-only.docx")
        self.assertEqual(manifest["sources"][0]["repository_path"], RAW_PATH)
        self.assertNotIn("file://", json.dumps(manifest).lower())

    def test_wiki_rel_source_matches_without_a_file(self) -> None:
        packet = copy.deepcopy(self.packet)
        packet["sources"][0]["path"] = WIKI_PATH
        self.assertEqual(validate_packet(packet, catalog=self.catalog), [])
        manifest = self.render(packet, "wiki-path.docx")
        self.assertEqual(manifest["sources"][0]["repository_path"], WIKI_PATH)
        self.assertTrue(
            manifest["sources"][0]["public_url"].endswith(f"/documents/{DOCUMENT_ID}/")
        )

    def test_validator_rejects_unsupported_fact_promotion(self) -> None:
        self.packet["evidence_matrix"][0]["basis"] = "unilateral_assertion"
        errors = validate_packet(self.packet, catalog=self.catalog)
        self.assertTrue(any("incompatible status/basis" in error for error in errors), errors)

    def test_validator_is_general_but_honors_packet_policy(self) -> None:
        self.assertEqual(validate_packet(self.packet, catalog=self.catalog), [])
        self.packet["metadata"]["process_number"] = "wrong"
        errors = validate_packet(self.packet, catalog=self.catalog)
        self.assertTrue(any("must remain '42/2026'" in error for error in errors), errors)

    def test_validator_checks_award_and_case_value_arithmetic(self) -> None:
        self.packet["decision_total"] = "0.50"
        self.packet["metadata"]["case_value"] = "2.00"
        errors = validate_packet(self.packet, catalog=self.catalog)
        self.assertTrue(any("award arithmetic mismatch" in error for error in errors))
        self.assertTrue(any("case value mismatch" in error for error in errors))

    def test_decision_total_mismatch_is_reported(self) -> None:
        self.packet["decision_total"] = "0.50"
        errors = validate_packet(self.packet, catalog=self.catalog)
        self.assertTrue(any("award arithmetic mismatch" in error for error in errors), errors)

    def test_not_proved_evidence_rejected_in_factos_provados(self) -> None:
        self.packet["sections"]["factos_provados"][0]["evidence_ids"] = ["u1"]
        errors = validate_packet(self.packet, catalog=self.catalog)
        self.assertTrue(
            any("relies on non-proved evidence u1" in error for error in errors),
            errors,
        )

    def test_sumario_requires_four_or_five_points(self) -> None:
        self.packet["sections"]["sumario"] = self.packet["sections"]["sumario"][:3]
        errors = validate_packet(self.packet, catalog=self.catalog)
        self.assertTrue(any("four or five" in error for error in errors), errors)

    def test_extra_paragraph_keys_are_ignored(self) -> None:
        for items in self.packet["sections"].values():
            for spec in items:
                spec["id"] = "para-extra-id"
                spec["confidence"] = "low"
                spec["note"] = "NOTE_SHOULD_NOT_RENDER"
        self.assertEqual(validate_packet(self.packet, catalog=self.catalog), [])
        docx = self.root / "extras.docx"
        render_docx(
            self.packet,
            template=TEMPLATE,
            out_docx=docx,
            catalog=self.catalog,
        )
        text = _document_text(docx)
        self.assertIn("Texto jurídico local.", text)
        self.assertNotIn("NOTE_SHOULD_NOT_RENDER", text)
        self.assertNotIn("para-extra-id", text)
        self.assertIn(WARNING, text)

    def test_load_local_catalog_tolerates_missing_open_sources(self) -> None:
        data = self.root / "data"
        data.mkdir()
        row = {
            "id": DOCUMENT_ID,
            "title": "Test authority",
            "rel_raw": RAW_PATH,
            "rel_source": WIKI_PATH,
        }
        (data / "catalog.jsonl").write_text(json.dumps(row) + "\n", encoding="utf-8")
        self.assertFalse((data / "open-sources-catalog.jsonl").exists())
        self.assertFalse((self.root / RAW_PATH).exists())
        catalog = load_local_catalog(self.root)
        self.assertEqual(
            catalog[DOCUMENT_ID],
            {"path": RAW_PATH, "title": "Test authority", "rel_source": WIKI_PATH},
        )
        self.assertEqual(validate_packet(self.packet, catalog=catalog), [])

    def test_load_local_catalog_merges_open_sources(self) -> None:
        data = self.root / "data"
        data.mkdir()
        open_id = "law_abcdefabcdefabcd"
        open_path = f"raw/documents/opensrc/2026/{open_id}.md"
        (data / "catalog.jsonl").write_text(
            json.dumps(
                {
                    "id": DOCUMENT_ID,
                    "title": "Test authority",
                    "rel_raw": RAW_PATH,
                    "rel_source": WIKI_PATH,
                }
            )
            + "\n",
            encoding="utf-8",
        )
        (data / "open-sources-catalog.jsonl").write_text(
            json.dumps(
                {
                    "id": open_id,
                    "title": "Open statute",
                    "rel_raw": open_path,
                    "rel_source": None,
                }
            )
            + "\n",
            encoding="utf-8",
        )
        catalog = load_local_catalog(self.root)
        self.assertEqual(catalog[open_id]["path"], open_path)
        self.assertEqual(catalog[open_id]["title"], "Open statute")
        self.assertIsNone(catalog[open_id]["rel_source"])
        self.assertIn(DOCUMENT_ID, catalog)

    def test_win32com_is_not_imported_at_module_level(self) -> None:
        source = (APP / "engine" / "render.py").read_text(encoding="utf-8")
        tree = ast.parse(source)
        for node in tree.body:
            if isinstance(node, ast.Import):
                names = [alias.name for alias in node.names]
            elif isinstance(node, ast.ImportFrom):
                names = [node.module or ""]
            else:
                continue
            for name in names:
                self.assertFalse(name.startswith("win32com"), name)


class EngineExampleTests(unittest.TestCase):
    def test_example_packet_validates_and_renders_minuta_marker(self) -> None:
        packet = json.loads(EXAMPLE_PACKET.read_text(encoding="utf-8"))
        catalog = load_local_catalog(ROOT)
        missing = [
            source["document_id"]
            for source in packet["sources"]
            if source["document_id"] not in catalog
        ]
        self.assertEqual(missing, [], f"example sources missing from catalog: {missing}")
        errors = validate_packet(packet, catalog=catalog)
        self.assertEqual(errors, [])
        with tempfile.TemporaryDirectory() as tmp:
            docx = Path(tmp) / "wizz.docx"
            manifest = render_docx(
                packet,
                template=TEMPLATE,
                out_docx=docx,
                catalog=catalog,
            )
            text = _document_text(docx)
            self.assertIn(WARNING, text)
            self.assertNotIn("file://", json.dumps(manifest).lower())
            self.assertEqual(validate_docx(docx, packet), [])
            self.assertEqual(validate_source_manifest(manifest, packet), [])
            for source in manifest["sources"]:
                document_id = source["document_id"]
                self.assertEqual(
                    source["public_url"],
                    f"https://dong-xuyong.github.io/law-wiki/documents/{document_id}/",
                )


class EngineExtractTests(unittest.TestCase):
    def test_extract_pdf_bytes_page_markers_and_digest(self) -> None:
        data = _text_pdf(
            [
                "Relato de factos suficientemente longo para extracao.",
                "Resposta processual igualmente suficientemente longa.",
            ]
        )
        extracted = extract_pdf_bytes(data, "case-record.pdf")
        text = extracted["text"]
        manifest = extracted["manifest"]
        self.assertIn("--- Página 1 ---", text)
        self.assertIn("--- Página 2 ---", text)
        self.assertIn("Relato de factos suficientemente longo para extracao.", text)
        self.assertEqual(manifest["pages"], 2)
        self.assertEqual(manifest["pages_with_text"], 2)
        self.assertEqual(manifest["empty_pages"], [])
        self.assertEqual(manifest["filename"], "case-record.pdf")
        self.assertEqual(manifest["sha256"], hashlib.sha256(data).hexdigest())
        self.assertEqual(manifest["extraction_method"], "pypdf")
        self.assertIn("not a Law Wiki legal authority", manifest["classification"])
        self.assertNotIn("\\", json.dumps(manifest))
        self.assertNotRegex(json.dumps(manifest), r"[A-Za-z]:[\\/]")

    def test_extracts_example_case_pdf(self) -> None:
        data = EXAMPLE_PDF.read_bytes()
        extracted = extract_pdf_bytes(data, EXAMPLE_PDF.name)
        self.assertIn("--- Página 1 ---", extracted["text"])
        self.assertIn("--- Página 2 ---", extracted["text"])
        self.assertIn("--- Página 3 ---", extracted["text"])
        self.assertEqual(extracted["manifest"]["pages"], 3)
        self.assertEqual(extracted["manifest"]["sha256"], hashlib.sha256(data).hexdigest())

    def test_rejects_image_only_pdf(self) -> None:
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        buffer = BytesIO()
        writer.write(buffer)
        with self.assertRaisesRegex(ValueError, "OCR is required"):
            extract_pdf_bytes(buffer.getvalue(), "scan.pdf")

    def test_rejects_encrypted_pdf(self) -> None:
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        writer.encrypt("secret")
        buffer = BytesIO()
        writer.write(buffer)
        with self.assertRaisesRegex(ValueError, "Encrypted"):
            extract_pdf_bytes(buffer.getvalue(), "locked.pdf")

    def test_rejects_pdf_with_no_pages(self) -> None:
        writer = PdfWriter()
        buffer = BytesIO()
        writer.write(buffer)
        with self.assertRaisesRegex(ValueError, "no pages"):
            extract_pdf_bytes(buffer.getvalue(), "empty.pdf")

    def test_rejects_non_pdf_filename(self) -> None:
        with self.assertRaisesRegex(ValueError, "must be a PDF"):
            extract_pdf_bytes(b"%PDF-1.7", "notes.txt")


if __name__ == "__main__":
    unittest.main()
